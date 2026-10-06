import io
import re
from pathlib import Path

import pytest
from docx import Document

from app import create_app
from db.connection import set_database_path
from db.migrate_runner import apply_migrations
from db.repositories import projects_repo
from services import evidence_service, research_task_service, tools


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    set_database_path(str(tmp_path / "test.db"))
    apply_migrations()
    flask_app = create_app()
    flask_app.config["TESTING"] = True
    with flask_app.test_client() as c:
        with flask_app.app_context():
            c.post("/api/projects", json={"name": "P"})
            yield c
    set_database_path(None)


def _seed(project="P", facts=None):
    pid = projects_repo.get_or_create_id(project)
    card = research_task_service.create_task(pid, "4", "Fees", research_method="TARGETED_WEB")
    facts = facts or [
        {"claim": "Deakin charges $3,000 per unit", "source_url": "https://deakin.example/fees",
         "source_title": "Deakin fees", "as_of": "2026"},
        {"claim": "Monash runs three intakes", "quote": "three intakes a year"},
    ]
    data = tools.run_tool("record_research_facts", "supervisor", project, {"task_id": card, "facts": facts,
                          "conclusions": [{"text": "Fees and intakes differ", "fact_numbers": [1, 2]}]}).data
    return card, data


def test_list_groups_facts_and_conclusions_by_phase(client):
    card, data = _seed()
    body = client.get("/api/evidence?project=P").get_json()
    assert body["success"] is True and set(body["phases"]) == {"4"}
    phase = body["phases"]["4"]
    assert [f["claim"] for f in phase["facts"]] == ["Deakin charges $3,000 per unit", "Monash runs three intakes"]
    assert phase["facts"][0]["source"] == {"url": "https://deakin.example/fees", "title": "Deakin fees",
                                           "publisher": "", "published_date": ""}
    assert (phase["facts"][0]["card_title"], phase["facts"][0]["as_of"]) == ("Fees", "2026")
    assert phase["facts"][1]["source"] is None and phase["facts"][1]["quote"] == "three intakes a year"
    (conclusion,) = phase["conclusions"]
    assert conclusion["text"] == "Fees and intakes differ"
    assert conclusion["fact_ids"] == data["fact_ids"] and conclusion["rejected_fact_count"] == 0


def test_list_for_one_phase(client):
    _seed()
    assert set(client.get("/api/evidence?project=P&phase=4").get_json()["phases"]) == {"4"}
    assert client.get("/api/evidence?project=P&phase=2").get_json()["phases"] == {}
    assert client.get("/api/evidence?project=P&phase=9").status_code == 400


def test_a_project_with_no_facts_lists_nothing(client):
    assert client.get("/api/evidence?project=P").get_json()["phases"] == {}


def test_search(client):
    _seed()
    body = client.get("/api/evidence/search?project=P&q=deakin").get_json()
    assert [f["claim"] for f in body["facts"]] == ["Deakin charges $3,000 per unit"]
    assert client.get("/api/evidence/search?project=P&q=" + "x" * 500).status_code == 200
    assert client.get("/api/evidence/search?project=P&q=deakin&phase=2").get_json()["facts"] == []
    assert client.get("/api/evidence/search?project=P&q=deakin&phase=9").status_code == 400


def test_reject_and_restore_a_fact(client):
    _, data = _seed()
    fact_id = data["fact_ids"][0]
    resp = client.post(f"/api/evidence/fact/{fact_id}/reject", json={"project": "P"})
    assert resp.status_code == 200 and resp.get_json() == {"success": True, "id": fact_id, "status": "rejected"}
    phase = client.get("/api/evidence?project=P").get_json()["phases"]["4"]
    assert phase["facts"][0]["status"] == "rejected"
    assert phase["conclusions"][0]["rejected_fact_count"] == 1
    client.post(f"/api/evidence/fact/{fact_id}/restore", json={"project": "P"})
    assert client.get("/api/evidence?project=P").get_json()["phases"]["4"]["facts"][0]["status"] == "active"


def test_reject_a_conclusion(client):
    _, data = _seed()
    conclusion_id = data["conclusion_ids"][0]
    assert client.post(f"/api/evidence/conclusion/{conclusion_id}/reject", json={"project": "P"}).get_json()["status"] == "rejected"


@pytest.mark.parametrize("url", ["/api/evidence/table/1/reject", "/api/evidence/fact/1/delete", "/api/evidence/fact/99999/reject"])
def test_unknown_things_are_404(client, url):
    _seed()
    assert client.post(url, json={"project": "P"}).status_code == 404


OTHER_FACTS = [{"claim": "Torrens charges $2,500 per unit"}, {"claim": "Federation runs two intakes"}]


def _other_project(client):
    """Project Q with its own facts and conclusion, and project P with its own."""
    client.post("/api/projects", json={"name": "Q"})
    _, mine = _seed("P")
    _, theirs = _seed("Q", OTHER_FACTS)
    return mine, theirs


def _q_rows(client):
    phase = client.get("/api/evidence?project=Q").get_json()["phases"]["4"]
    return {f["id"]: f["status"] for f in phase["facts"]}, {c["id"]: c["status"] for c in phase["conclusions"]}


@pytest.mark.parametrize("action", ["reject", "restore"])
def test_a_fact_or_conclusion_of_another_project_cannot_be_changed(client, action):
    _, theirs = _other_project(client)
    facts_before, conclusions_before = _q_rows(client)
    for kind, item in (("fact", theirs["fact_ids"][0]), ("conclusion", theirs["conclusion_ids"][0])):
        resp = client.post(f"/api/evidence/{kind}/{item}/{action}", json={"project": "P"})
        assert resp.status_code == 404 and resp.get_json()["success"] is False
    assert _q_rows(client) == (facts_before, conclusions_before)
    assert set(facts_before.values()) == {"active"}


def test_the_other_project_changes_only_through_its_own_name(client):
    _, theirs = _other_project(client)
    fact_id = theirs["fact_ids"][0]
    assert client.post(f"/api/evidence/fact/{fact_id}/reject", json={"project": "Q"}).status_code == 200
    assert _q_rows(client)[0][fact_id] == "rejected"


def test_list_and_search_never_include_another_projects_facts(client):
    mine, theirs = _other_project(client)
    listed = client.get("/api/evidence?project=P").get_json()["phases"]["4"]
    assert [f["id"] for f in listed["facts"]] == mine["fact_ids"]
    assert [c["id"] for c in listed["conclusions"]] == mine["conclusion_ids"]
    assert not set(theirs["fact_ids"]) & {f["id"] for f in listed["facts"]}
    assert client.get("/api/evidence/search?project=P&q=torrens").get_json()["facts"] == []
    assert client.get("/api/evidence/search?project=P&q=intakes").get_json()["facts"][0]["claim"] == "Monash runs three intakes"
    found = client.get("/api/evidence/search?project=Q&q=intakes").get_json()["facts"]
    assert [f["claim"] for f in found] == ["Federation runs two intakes"]


def test_no_project_is_400(client):
    responses = [
        client.get("/api/evidence?project=Nope"),
        client.get("/api/evidence/search?project=Nope&q=x"),
        client.post("/api/evidence/fact/1/reject", json={"project": "Nope"}),
    ]
    for response in responses:
        assert response.status_code == 400
        assert response.get_json() == {"success": False, "error": "No project selected."}


def test_the_page_loads_the_evidence_script_and_search_box(client):
    html = client.get("/").get_data(as_text=True)
    assert 'src="/static/evidence.js' in html and 'id="evidence-search-input"' in html


def test_facts_are_not_drawn_under_a_past_insights_version():
    # app.js cannot run under Node, so this checks the gate in the source: the slot that evidence.js fills
    # is only written when the tab is not showing a frozen history entry (facts are not versioned).
    source = (Path(__file__).resolve().parents[1] / "static" / "app.js").read_text(encoding="utf-8")
    assert re.search(r"\$\{viewingHistoricalVersion \? '' : `<div class=\"ev-slot\" data-evidence-phase=", source)


# ---- the facts briefing for Generate ----

def test_the_brief_for_a_phase(client):
    _, data = _seed()
    first, second = data["fact_ids"]
    body = client.get("/api/evidence/brief?project=P&phase=4").get_json()
    assert body["success"] is True
    brief = body["brief"]
    assert (brief["fact_count"], brief["conclusion_count"], brief["shown_facts"]) == (2, 1, 2)
    assert brief["rejected_count"] == 0
    assert brief["text"].startswith("# CHECKED FACTS FOR PHASE 4 (newest first)\n"
                                    f"[F{second}] Monash runs three intakes\n"
                                    f"[F{first}] Deakin charges $3,000 per unit — Deakin fees, as of 2026\n")
    assert client.get("/api/evidence/brief?project=P&phase=2").get_json()["brief"]["fact_count"] == 0


def test_the_brief_reports_how_many_facts_are_rejected(client):
    _, data = _seed()
    first, _second = data["fact_ids"]
    client.post(f"/api/evidence/fact/{first}/reject", json={"project": "P"})
    brief = client.get("/api/evidence/brief?project=P&phase=4").get_json()["brief"]
    assert (brief["fact_count"], brief["rejected_count"]) == (1, 1)
    assert "# REJECTED — DO NOT USE\n- Deakin charges $3,000 per unit\n" in brief["text"]
    assert client.get("/api/evidence/brief?project=P&phase=2").get_json()["brief"]["rejected_count"] == 0


@pytest.mark.parametrize("query", ["", "&phase=", "&phase=9", "&phase=x"])
def test_the_brief_needs_a_phase_from_1_to_7(client, query):
    resp = client.get(f"/api/evidence/brief?project=P{query}")
    assert resp.status_code == 400
    assert resp.get_json() == {"success": False, "error": "Choose a phase from 1 to 7."}


def test_the_brief_with_no_project_is_400(client):
    resp = client.get("/api/evidence/brief?project=Nope&phase=4")
    assert resp.status_code == 400 and resp.get_json() == {"success": False, "error": "No project selected."}


def test_the_brief_never_includes_another_projects_facts(client):
    _other_project(client)
    text = client.get("/api/evidence/brief?project=P&phase=4").get_json()["brief"]["text"]
    assert "Deakin" in text and "Torrens" not in text and "Federation" not in text


# ---- the Word report ----

def test_the_word_report_lists_the_sources_of_this_projects_facts_only(client):
    mine, theirs = _other_project(client)
    deakin, monash = mine["fact_ids"]
    torrens = theirs["fact_ids"][0]
    insights = {"competitors": [], "phases": {
        "4": {"title": "Product Features", "summary": f"Deakin is dear [F{deakin}]. Torrens [F{torrens}]. [F{monash}]",
              "confidence": "medium"},
        "5": {"title": "Academic Content", "summary": "From the reports only.", "confidence": "low"},
    }}
    resp = client.post("/api/insights/report", json={"insights": insights, "project_name": "P"})
    assert resp.status_code == 200
    texts = [p.text for p in Document(io.BytesIO(resp.data)).paragraphs]
    start = texts.index("Cited sources")
    assert texts[start + 1:start + 4] == [
        f"F{deakin} — Deakin fees (2026) — https://deakin.example/fees",
        f"F{torrens} — (not found)",
        f"F{monash} — Research report: Fees",
    ]
    assert texts.count("Cited sources") == 1 and evidence_service.NOT_CITED_NOTE in texts


def test_the_word_report_still_downloads_when_the_citation_lookup_fails(client, monkeypatch):
    mine, _ = _other_project(client)
    deakin = mine["fact_ids"][0]

    def broken(project_name):
        raise RuntimeError("database is locked")

    monkeypatch.setattr(evidence_service, "citation_index", broken)
    insights = {"competitors": [], "phases": {
        "4": {"title": "Product Features", "summary": f"Deakin is dear [F{deakin}].", "confidence": "medium"}}}
    resp = client.post("/api/insights/report", json={"insights": insights, "project_name": "P"})
    assert resp.status_code == 200
    texts = [p.text for p in Document(io.BytesIO(resp.data)).paragraphs]
    assert f"F{deakin} — (not found)" in texts


def test_the_word_report_downloads_when_stored_text_has_control_characters(client):
    junk = "\x00\x01\x08\x1b\ufffe\uffff"
    _, data = _seed("P", [{"claim": "Deakin charges more" + junk, "source_url": "https://deakin.example/fees",
                           "source_title": "a\x00b", "publisher": "c\x1bd" + junk, "as_of": "2026" + junk}])
    fact_id = data["fact_ids"][0]
    conclusion_id = data["conclusion_ids"][0]
    insights = {"competitors": [], "phases": {
        "4": {"title": "Product Features", "summary": f"Fees [F{fact_id}] [C{conclusion_id}].", "confidence": "medium"}}}
    resp = client.post("/api/insights/report", json={"insights": insights, "project_name": "P"})
    assert resp.status_code == 200
    texts = [p.text for p in Document(io.BytesIO(resp.data)).paragraphs]
    assert f"F{fact_id} — cd — ab (2026) — https://deakin.example/fees" in texts
    brief = client.get("/api/evidence/brief?project=P&phase=4").get_json()["brief"]["text"]
    assert not re.search("[\x00-\x08\x0e-\x1f\ufffe\uffff]", brief)


# ---- app.js glue (app.js cannot run under Node, so these check its source) ----

def _static(name):
    return (Path(__file__).resolve().parents[1] / "static" / name).read_text(encoding="utf-8")


def _between(source, start, end):
    return source[source.index(start):source.index(end, source.index(start))]


def test_generate_adds_the_briefing_after_the_guidance_and_before_the_users_instructions():
    body = _between(_static("app.js"), "async function generatePhaseInsight(", "async function handleInsightResponse(")
    added = body.index("prompt += await evidencePhaseBriefAddition(briefProject, phaseKey);")
    assert body.index("CRITICAL INSTRUCTIONS:") < added < body.index("USER INSTRUCTIONS:")
    assert "typeof evidencePhaseBriefAddition === 'function'" in body       # an old cached evidence.js adds nothing
    switched = body.index("if (currentProject !== briefProject) {")        # a project switch while it loads stops
    assert added < switched


def test_generate_checks_for_a_running_generation_before_and_after_waiting_for_the_briefing():
    body = _between(_static("app.js"), "async function generatePhaseInsight(", "async function handleInsightResponse(")
    added = body.index("prompt += await evidencePhaseBriefAddition(briefProject, phaseKey);")
    check = "if (refuseIfRunning()) return;"
    assert body.count(check) == 2
    first = body.index(check)
    second = body.index(check, first + 1)
    assert first < added < second                                           # before the wait, and again once it is over
    assert body.index("if (currentProject !== briefProject) {") < second
    assert body.count("An insight generation is already running for this project.") == 1   # one message, one helper
    assert body.count("getActiveGenerationForProject(currentProject)") == 1


def test_phase_7_keeps_the_markers_word_for_word_with_the_options_report():
    source = _static("app.js")
    assert f"const KEEP_CITATION_MARKERS = '{evidence_service.KEEP_MARKERS_INSTRUCTION}';" in source
    body = _between(source, "function buildPriorPhaseContextForPhase7(", "function inferFieldConfidence(")
    assert "if (sections.length === 0) return '';" in body                   # still empty when there is nothing
    assert r"return `${KEEP_CITATION_MARKERS}\n\n${clipped}`;" in body
